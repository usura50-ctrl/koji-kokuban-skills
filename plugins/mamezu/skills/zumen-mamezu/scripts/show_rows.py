"""作ったExcelの行を、画面に出す（図面と見比べて、値を確認するため）。
使い方：python show_rows.py <Excel> [まとまり] [最初の番号=1] [件数=10]"""
import sys

from openpyxl import load_workbook

path = sys.argv[1]
group = sys.argv[2] if len(sys.argv) > 2 else ''
start = int(sys.argv[3]) if len(sys.argv) > 3 else 1
count = int(sys.argv[4]) if len(sys.argv) > 4 else 10
ws = load_workbook(path)['撮影予定']
rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] and (not group or r[9] == group)]
print(f'{group or "全て"}：{len(rows)}行')
for r in rows[start - 1:start - 1 + count]:
    print(f'--- No.{r[0]} [{r[9]}] {r[1]} / 豆図：{r[10] or "なし"} {r[11] or ""}')
    print(r[6])
