"""図面PDFの一式から、全ての一覧表を自動で見つけ、豆図（PNG）と黒板Excelを作る（実験用）。
使い方：python run_project.py <PDF> <出力フォルダ> [工種]"""
import os
import re
import shutil
import sys
from collections import Counter

import pymupdf
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image

from cellcrop import cell_crop
from extract import compose, find_habadome, floor_key, natural, records, slab_compose, slab_records
from listkit import Lists, nfkc, squash

PDF = sys.argv[1]
OUT = sys.argv[2]
KOSHU = sys.argv[3] if len(sys.argv) > 3 else '鉄筋工事'
GROUPS = ['基礎梁', '柱', '梁', 'スラブ', '壁']

L = Lists(PDF)
habadome = find_habadome(L)
print('幅止め筋：', habadome)
shutil.rmtree(OUT, ignore_errors=True)
os.makedirs(OUT)
report = []  # 人が確認すること


def group_of(tb):
    """表の上の題名（地中梁リスト・柱リスト…）から、まとまりを決める。"""
    pg = L.page(tb.page)
    t = squash(''.join(w for r, w in sorted(pg.words_in(tb.x0 - 5, tb.head[0] - 45, tb.x0 + 330, tb.head[0] - 1), key=lambda a: (round(a[0].y0 / 5), a[0].x0))))
    if 'スラブ' in t:
        return 'スラブ', t
    if '地中' in t and '梁' in t:
        return '基礎梁', t
    if '大梁' in t or '小梁' in t or '梁リスト' in t:
        return '梁', t
    if '柱' in t:
        return '柱', t
    if '壁' in t:
        return '壁', t
    return None, t


def sheet_no(pno):
    ws = L.doc[pno].get_text('words')
    cand = [nfkc(w[4]) for w in ws if re.fullmatch(r'S-\d+', nfkc(w[4]))]
    return cand[-1] if cand else f'p{pno + 1}'


items = []
for pno in range(len(L.doc)):
    for tb in L.find_tables(pno):
        group, title = group_of(tb)
        if group is None:
            if tb.kind != 'slab' or True:
                report.append(f'{sheet_no(pno)}：題名が読めない表を飛ばしました（{title[:20] or "題名なし"}）')
            continue
        if tb.kind == 'slab' or group == 'スラブ':
            try:
                recs = slab_records(L, tb)
            except ValueError:
                recs = []
            for r in recs:
                items.append(dict(group=group, name=r['sym'], naiyou=slab_compose(r), order=(0, natural(r['sym']), 0), image=None, sheet=sheet_no(pno), kind='標準', width=''))
            if not recs:
                report.append(f'{sheet_no(pno)}：スラブの表（{title[:20]}）から符号を読めませんでした')
            continue
        recs = records(L, tb)
        # 符号があるのに、値が空の列（「S-26参照」など）を知らせる
        got = {r['sym'] for r in recs}
        for a, b in tb.cols:
            sym = nfkc(squash(''.join(w for rr, w in L.page(pno).words_in(a, tb.head[0], b, tb.head[1]))))
            if sym and sym not in got:
                report.append(f'{sheet_no(pno)}：{sym} は、表に値がないので、黒板を作りませんでした（他の図を参照？）')
        for r in recs:
            name = r['sym'] + (f"_{r['floor']}階" if r['floor'] else '') + r['pos']
            items.append(dict(group=group, name=name, naiyou=compose(r, tb.kind, habadome if tb.kind in ('beam', 'wall') else None),
                              order=(floor_key(r['floor']), natural(r['sym']), r['cell'][0]), page=pno, cell=r['cell'], sheet=sheet_no(pno)))

# ---- 画像
ROW_UNIT, PAD_UNIT = 0.075, 0.18


def board_for(w, h):
    if h <= w:
        return '豆図付き', '普通'
    pad = ROW_UNIT * PAD_UNIT
    need = (8 * ROW_UNIT - 2 * pad) * (w / h) + 2 * pad
    return '豆図付き（上部拡大）', ('狭い' if need <= 0.40 else ('普通' if need <= 0.55 else '広い'))


names = Counter(i['name'] for i in items if i.get('cell'))
for n, c in names.items():
    if c > 1:
        report.append(f'同じ名前の黒板が{c}件あります（図面の間違いの可能性）：{n}')
doc = pymupdf.open(PDF)
for it in items:
    if not it.get('cell'):
        continue
    path = os.path.join(OUT, it['group'], it['name'] + '.png')
    try:
        w, h = cell_crop(doc, it['page'], it['cell'], path)
    except ValueError:
        report.append(f"{it['sheet']}：{it['name']} は、断面の欄に図がなかったので、画像を作りませんでした")
        it['image'] = None
        it['kind'], it['width'] = '標準', ''
        continue
    it['image'] = it['name'] + '.png'
    it['kind'], it['width'] = board_for(w, h)

for it in items:
    if re.search(r'：-[DKU]', it['naiyou']):
        report.append(f"{it['sheet']}：{it['name']} の値の頭に記号（□など）が付いていません。図の記号を確認してください：{it['naiyou'].splitlines()[-1] if False else [l for l in it['naiyou'].splitlines() if re.search('：-[DKU]', l)][0]}")
items.sort(key=lambda a: (GROUPS.index(a['group']), a['order']))

wb = Workbook()
sh = wb.active
sh.title = '撮影予定'
H = ['No.', '黒板の種類', '工種', '種別', '細別', '測点', '作業内容', '必要枚数', '撮影予定日', 'まとまり', '豆図ファイル', '豆図領域幅']
W = [6, 20, 14, 12, 12, 12, 40, 10, 14, 14, 26, 12]
fill = PatternFill('solid', fgColor='FFF26B0F')
thin = Side(style='thin', color='FFC9CDD3')
for i, (h, w) in enumerate(zip(H, W), 1):
    c = sh.cell(1, i, h)
    c.fill = fill
    c.font = Font(bold=True, color='FFFFFFFF')
    c.alignment = Alignment(horizontal='center', vertical='center')
    sh.column_dimensions[c.column_letter].width = w
sh.freeze_panes = 'A2'
for n, it in enumerate(items, 1):
    sh.append([n, it['kind'], KOSHU, '', '', '', it['naiyou'], 1, None, it['group'], it['image'] or '', it['width']])
    sh.row_dimensions[n + 1].height = 15 * (it['naiyou'].count('\n') + 1) + 6
for rr in range(2, max(300, len(items) + 100) + 1):
    for c in range(1, len(H) + 1):
        sh.cell(rr, c).border = Border(top=thin, bottom=thin, left=thin, right=thin)
    sh.cell(rr, 7).alignment = Alignment(wrap_text=True, vertical='top')
    sh.cell(rr, 9).number_format = 'yyyy/mm/dd'
wb.save(os.path.join(OUT, '黒板_基礎梁・柱・梁・スラブ・壁.xlsx'))

with open(os.path.join(OUT, '確認してください.txt'), 'w', encoding='utf-8') as f:
    f.write('\n'.join(report) if report else '確認が必要な点はありません。')
print(len(items), '行', Counter(i['group'] for i in items))
print(Counter((i['kind'], i['width']) for i in items))
print('\n'.join(report))
